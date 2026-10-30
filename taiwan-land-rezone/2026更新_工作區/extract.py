# -*- coding: utf-8 -*-
import openpyxl, json, os

SRC = r"D:\_WWW_325_private\MegaMenu一目了然\________325_台灣22縣市重劃區megamenu___________\____台灣22縣市重劃區_詳細資料庫_線上整合版-------------------------.xlsx"
WORK = r"D:\_WWW_325_private\MegaMenu一目了然\________325_台灣22縣市重劃區megamenu___________\2026更新_工作區"

TARGETS = {"桃園市", "新竹市", "新竹縣", "苗栗縣"}

wb = openpyxl.load_workbook(SRC, read_only=True, data_only=True)
print("SHEETS:", wb.sheetnames)

# --- 建案總表 ---
ws = wb["建案總表"]
rows = list(ws.iter_rows(min_row=1, max_row=643, values_only=True))
header = rows[0]
print("HEADER:", header)
data = []
for r in rows[1:]:
    if r[1] in TARGETS:
        data.append(r)
print("TARGET ROW COUNT:", len(data))
# count by city
from collections import Counter
print(Counter([r[1] for r in data]))

# dump to json
out = []
for r in data:
    out.append({
        "序": r[0], "縣市": r[1], "縣市編碼": r[2], "重劃區名稱": r[3],
        "開發類型": r[4], "建案名稱": r[5], "建商": r[6], "戶數": r[7],
        "坪數區間": r[8], "單價區間": r[9], "推出年": r[10], "狀態": r[11],
        "資料來源等級": r[12], "資料來源網址": r[13],
    })
with open(os.path.join(WORK, "target_rows.json"), "w", encoding="utf-8") as f:
    json.dump(out, f, ensure_ascii=False, indent=1)

# --- 全站重劃區清單 ---
ws2 = wb["全站重劃區清單"]
rows2 = list(ws2.iter_rows(values_only=True))
print("=== 全站重劃區清單 header:", rows2[0])
districts = []
for r in rows2[1:]:
    if r and r[0] in TARGETS or (len(r)>1 and r[1] in TARGETS):
        districts.append(r)
print("DISTRICTS target count:", len(districts))
with open(os.path.join(WORK, "districts.json"), "w", encoding="utf-8") as f:
    json.dump([list(r) for r in districts], f, ensure_ascii=False, indent=1)

wb.close()
